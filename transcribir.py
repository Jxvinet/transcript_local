"""Transcribe video/audio en local con faster-whisper (modelos OpenAI Whisper)."""
from __future__ import annotations

import argparse
import configparser
import os
import shutil
import ssl
import sys
from pathlib import Path
from typing import Callable

ROOT = Path(__file__).resolve().parent
AJUSTES = ROOT / "ajustes.ini"
AUDIO_EXT = {
    ".mp4", ".mkv", ".mov", ".avi", ".webm",
    ".mp3", ".wav", ".m4a", ".wma", ".flac", ".ogg",
}


def load_ajustes(path: Path = AJUSTES) -> dict:
    cfg = configparser.ConfigParser()
    defaults = {
        "modelo": "small",
        "idioma": "es",
        "carpeta": "",
        "formatos": "docx",
        "dispositivo": "cpu",
        "calculo": "int8",
        "ssl_inseguro": "0",
    }
    if path.exists():
        # ajustes.ini sin seccion [..] -> lo leemos como DEFAULT-like
        text = path.read_text(encoding="utf-8")
        cfg.read_string("[ajustes]\n" + text)
        sec = cfg["ajustes"]
        for key in defaults:
            if key not in sec:
                continue
            # idioma/carpeta vacios son validos (auto / Escritorio\Transcripciones)
            raw = str(sec[key]).strip()
            if key in ("idioma", "carpeta") or raw != "":
                defaults[key] = raw
    return defaults


def apply_ssl_workaround(enabled: bool) -> None:
    """Solo si la red corporativa rompe Hugging Face. No activar en publico por defecto."""
    if not enabled:
        return
    ssl._create_default_https_context = ssl._create_unverified_context
    os.environ.setdefault("PYTHONHTTPSVERIFY", "0")
    os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")


def find_ffmpeg() -> Path | None:
    """Busca ffmpeg en PATH, FFMPEG_PATH o instalaciones tipicas de WinGet."""
    env = os.environ.get("FFMPEG_PATH", "").strip()
    if env:
        p = Path(env)
        if p.is_file():
            return p
        if (p / "ffmpeg.exe").is_file():
            return p / "ffmpeg.exe"

    which = shutil.which("ffmpeg")
    if which:
        return Path(which)

    winget = Path.home() / "AppData/Local/Microsoft/WinGet/Packages"
    if winget.is_dir():
        matches = sorted(winget.glob("Gyan.FFmpeg*/**/ffmpeg.exe"), reverse=True)
        if matches:
            return matches[0]

    for base in (
        Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "ffmpeg" / "bin",
        Path(os.environ.get("LOCALAPPDATA", "")) / "ffmpeg" / "bin",
    ):
        cand = base / "ffmpeg.exe"
        if cand.is_file():
            return cand
    return None


def ensure_ffmpeg_on_path() -> Path | None:
    exe = find_ffmpeg()
    if exe is None:
        return None
    bin_dir = str(exe.parent)
    path = os.environ.get("PATH", "")
    if bin_dir.lower() not in path.lower():
        os.environ["PATH"] = bin_dir + os.pathsep + path
    return exe


def default_out_dir(ajustes: dict | None = None) -> Path:
    ajustes = ajustes or load_ajustes()
    raw = (ajustes.get("carpeta") or "").strip()
    if raw:
        return Path(os.path.expandvars(os.path.expanduser(raw)))
    return Path.home() / "Desktop" / "Transcripciones"


def fmt_srt(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = seconds % 60
    return f"{h:02d}:{m:02d}:{s:06.3f}".replace(".", ",")


def fmt_txt(seconds: float) -> str:
    return fmt_srt(seconds).replace(",", ".")


_model = None
_model_key: tuple | None = None


def get_model(name: str, device: str = "cpu", compute_type: str = "int8"):
    global _model, _model_key
    key = (name, device, compute_type)
    if _model is not None and _model_key == key:
        return _model
    from faster_whisper import WhisperModel

    try:
        _model = WhisperModel(name, device=device, compute_type=compute_type, local_files_only=True)
    except Exception:
        _model = WhisperModel(name, device=device, compute_type=compute_type)
    _model_key = key
    return _model


def transcribe_file(
    src: Path,
    modelo: str = "small",
    idioma: str = "es",
    device: str = "cpu",
    compute_type: str = "int8",
    on_segment: Callable[[float, float, str], None] | None = None,
) -> list[tuple[float, float, str]]:
    ensure_ffmpeg_on_path()
    model = get_model(modelo, device=device, compute_type=compute_type)
    kwargs: dict = {
        "vad_filter": True,
        "beam_size": 5,
    }
    if idioma:
        kwargs["language"] = idioma
    segments, _info = model.transcribe(str(src), **kwargs)
    rows: list[tuple[float, float, str]] = []
    for seg in segments:
        text = (seg.text or "").strip()
        if not text:
            continue
        rows.append((seg.start, seg.end, text))
        if on_segment:
            on_segment(seg.start, seg.end, text)
    return rows


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    n = 2
    while True:
        candidate = path.with_name(f"{stem}_{n}{suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def write_docx(dest: Path, src: Path, rows: list[tuple[float, float, str]]) -> Path:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    dest = unique_path(dest)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    title = doc.add_heading(src.stem, level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    meta = doc.add_paragraph()
    r = meta.add_run(f"Archivo original: {src.name}")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if not rows:
        doc.add_paragraph("(Sin texto detectado.)")
    else:
        for start, end, text in rows:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            ts = p.add_run(f"{fmt_txt(start)} – {fmt_txt(end)}  ")
            ts.bold = True
            ts.font.size = Pt(10)
            ts.font.color.rgb = RGBColor(0xB6, 0x75, 0x2E)
            body = p.add_run(text)
            body.font.size = Pt(12)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def write_txt(dest: Path, rows: list[tuple[float, float, str]]) -> Path:
    dest = unique_path(dest)
    lines = [f"[{fmt_txt(s)} --> {fmt_txt(e)}] {t}" for s, e, t in rows]
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    return dest


def write_srt(dest: Path, rows: list[tuple[float, float, str]]) -> Path:
    dest = unique_path(dest)
    blocks = [
        f"{i}\n{fmt_srt(s)} --> {fmt_srt(e)}\n{t}\n"
        for i, (s, e, t) in enumerate(rows, 1)
    ]
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(blocks) + ("\n" if blocks else ""), encoding="utf-8")
    return dest


def save_outputs(
    out_dir: Path,
    src: Path,
    rows: list[tuple[float, float, str]],
    formatos: list[str],
) -> list[Path]:
    written: list[Path] = []
    for fmt in formatos:
        f = fmt.strip().lower()
        if f == "docx":
            written.append(write_docx(out_dir / f"{src.stem}.docx", src, rows))
        elif f == "txt":
            written.append(write_txt(out_dir / f"{src.stem}.txt", rows))
        elif f == "srt":
            written.append(write_srt(out_dir / f"{src.stem}.srt", rows))
    return written


def main() -> int:
    ajustes = load_ajustes()
    apply_ssl_workaround(ajustes.get("ssl_inseguro", "0") in ("1", "true", "True", "yes"))

    parser = argparse.ArgumentParser(description="Transcribir reunion (Whisper, local)")
    parser.add_argument("archivo", help="Ruta al video o audio")
    parser.add_argument("--modelo", default=ajustes["modelo"], help="tiny|base|small|medium|large-v2|large-v3")
    parser.add_argument("--idioma", default=ajustes["idioma"], help="es, en, ... (vacio = auto)")
    parser.add_argument("--formatos", default=ajustes["formatos"], help="docx,txt,srt")
    parser.add_argument("--carpeta", default=ajustes["carpeta"], help="Carpeta de salida")
    args = parser.parse_args()

    src = Path(args.archivo)
    if not src.exists():
        print(f"No existe el archivo: {src}", file=sys.stderr)
        return 1
    if src.suffix.lower() not in AUDIO_EXT:
        print(f"Formato no reconocido: {src.suffix}", file=sys.stderr)
        return 1

    if ensure_ffmpeg_on_path() is None:
        print("No encuentro ffmpeg. Instálalo (winget install Gyan.FFmpeg) o añádelo al PATH.", file=sys.stderr)
        return 1

    out_dir = Path(args.carpeta) if args.carpeta.strip() else default_out_dir(ajustes)
    formatos = [x.strip() for x in args.formatos.split(",") if x.strip()]
    if not formatos:
        formatos = ["docx"]

    print(f"Archivo : {src}")
    print(f"Modelo  : {args.modelo} ({ajustes['dispositivo']}, idioma={args.idioma or 'auto'})")
    print(f"Salida  : {out_dir}")
    print("Cargando modelo...", flush=True)

    def on_segment(start, end, text):
        print(f"[{fmt_txt(start)} --> {fmt_txt(end)}] {text}", flush=True)

    rows = transcribe_file(
        src,
        modelo=args.modelo,
        idioma=args.idioma,
        device=ajustes["dispositivo"],
        compute_type=ajustes["calculo"],
        on_segment=on_segment,
    )
    written = save_outputs(out_dir, src, rows, formatos)
    print()
    for p in written:
        print(f"Listo: {p}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
